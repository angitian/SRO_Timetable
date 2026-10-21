"""เปลี่ยนชื่อครูจริงในไฟล์ตัวอย่าง 2 ไฟล์ → ชื่อไทยสมมุติ (one-off utility)

อ่านไฟล์ต้นฉบับ (ไม่ถูกแก้) แล้วสร้างไฟล์ตัวอย่างที่ไม่ระบุชื่อจริง
สำหรับแจกในแอป frontend/public/samples/

การใช้งาน:
    cd scripts
    python anonymize_sample_teachers.py

หรือระบุ path เอง:
    python anonymize_sample_teachers.py "<src_course.xlsx>" "<src_activities.xlsx>" "<out_dir>"

หมายเหตุ:
- ใช้ตารางจับคู่ชุดเดียวกับทั้ง 2 ไฟล์ → คู่ตัวอย่างยังใช้ร่วมกันได้
  (ชื่อครูในไฟล์คาบพิเศษ resolve กับไฟล์ข้อมูลการเรียนได้ครบ)
- ชื่อที่เป็น ASCII ล้วน (ครูต่างชาติ) → Mr./Ms. สมมุติ
- ชื่อไทย → อ. + ชื่อ-นามสกุลสมมุติ (สไตล์เดียวกับ MIN_TEMPLATE_EXAMPLES)
- แก้เฉพาะคอลัมน์ชื่อครู — โครงสร้าง/ค่าอื่นคงเดิมทุกอย่าง
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

# ── ชื่อสมมุติ ──────────────────────────────────────────────────────────────

THAI_PSEUDONYMS = [
    "อ.สมชาย ใจดี", "อ.สมหญิง รักเรียน", "อ.วิชัย สอนเก่ง",
    "อ.พิมพ์ชนก อารีย์", "อ.ธนกร ศรีสุข", "อ.นภัสสร วัฒนกุล",
    "อ.กิตติพงษ์ จันทร์เพ็ญ", "อ.สุพัตรา บุญนำ", "อ.อนุชา ทองแท้",
    "อ.วราภรณ์ ชูช่วย", "อ.ประเสริฐ พูนสุข", "อ.รัตนา แก้วมณี",
    "อ.สมบูรณ์ ศิริธรรม", "อ.จุฑารัตน์ เพิ่มพูน", "อ.ธวัชชัย พัฒนกิจ",
    "อ.มาลี บุญชู", "อ.นิรันดร์ สงบสุข", "อ.ปรีชา เลิศลาภ",
    "อ.ศิริพร วงศ์วิไล", "อ.วิไลวรรณ สุขใจ", "อ.ชัยวัฒน์ ตั้งตระกูล",
    "อ.อารียา จิตรงาม", "อ.สมศักดิ์ แสนยินดี", "อ.นิภาพร ปัญญาดี",
    "อ.ยุพดี ใจกว้าง", "อ.สมบัติ อุดมเดช", "อ.พรทิพย์ สายทอง",
    "อ.ทวีศักดิ์ พุฒซ้อน", "อ.อัมพร ชื่นใจ", "อ.สมเจตน์ ประเสริฐสุข",
    "อ.วาสนา คงดี", "อ.อุดม ศิริสมบูรณ์", "อ.เพ็ญศรี ผาสุข",
    "อ.สุรินทร์ เพียรดี", "อ.พัชรี แสงจันทร์", "อ.ประยุทธ์ กิตติศักดิ์",
    "อ.ธิดารัตน์ บุญมาก", "อ.สมโภชน์ สุขสันต์", "อ.จริญญา วงศ์ดี",
]

EN_PSEUDONYMS = [
    "Mr. Smith", "Ms. Johnson", "Mr. Brown", "Ms. Wilson", "Mr. Davis",
    "Ms. Taylor", "Mr. Miller", "Ms. Anderson",
]


def _is_ascii(s: str) -> bool:
    return all(ord(c) < 128 for c in s)


def _split_names(raw) -> list[str]:
    if raw is None:
        return []
    return [n.strip() for n in str(raw).split(",") if n.strip()]


def collect_names(course_path: Path, activities_path: Path) -> set[str]:
    """รวมชื่อครูจริงทั้งหมดจาก 2 ไฟล์

    course file: col0 ชื่อครู, col1 ครูผู้สอนร่วม
    activities file: col5 ครูผู้สอน (คั่นคอมม่า)
    """
    names: set[str] = set()

    wb = openpyxl.load_workbook(course_path, data_only=True)
    for row in list(wb.active.iter_rows(values_only=True))[1:]:
        if row[0]:
            names.add(str(row[0]).strip())
        if row[1]:
            names.add(str(row[1]).strip())

    wb = openpyxl.load_workbook(activities_path, data_only=True)
    for row in list(wb.active.iter_rows(values_only=True))[1:]:
        names.update(_split_names(row[5]))

    return names


def build_mapping(names: set[str]) -> dict[str, str]:
    """ชื่อจริง (เรียง) → ชื่อสมมุติ; ASCII → EN, ไทย → อ.xxx"""
    mapping: dict[str, str] = {}
    thai_i = en_i = 0
    for name in sorted(names):
        if _is_ascii(name):
            if en_i >= len(EN_PSEUDONYMS):
                raise RuntimeError(f"EN pseudonyms exhausted ({en_i})")
            mapping[name] = EN_PSEUDONYMS[en_i]
            en_i += 1
        else:
            if thai_i >= len(THAI_PSEUDONYMS):
                raise RuntimeError(f"Thai pseudonyms exhausted ({thai_i})")
            mapping[name] = THAI_PSEUDONYMS[thai_i]
            thai_i += 1
    return mapping


def anonymize_course(src: Path, dst: Path, mapping: dict[str, str]) -> None:
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        for col in (0, 1):
            cell = row[col]
            name = str(cell.value).strip() if cell.value is not None else ""
            if name and name in mapping:
                cell.value = mapping[name]
    wb.save(dst)


def anonymize_activities(src: Path, dst: Path, mapping: dict[str, str]) -> None:
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        cell = row[5]
        names = _split_names(cell.value)
        if names:
            cell.value = ", ".join(mapping.get(n, n) for n in names)
    wb.save(dst)


def main() -> int:
    root = Path(__file__).resolve().parent.parent
    course_src = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(
        r"C:\Users\Lenovo\Downloads\ลงตารางสอน2 15.10.69.xlsx"
    )
    activities_src = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(
        r"C:\Users\Lenovo\Downloads\คาบ.xlsx"
    )
    out_dir = Path(sys.argv[3]) if len(sys.argv) > 3 else root / "frontend" / "public" / "samples"
    out_dir.mkdir(parents=True, exist_ok=True)

    course_dst = out_dir / "ตัวอย่างข้อมูลการเรียน.xlsx"
    activities_dst = out_dir / "ตัวอย่างคาบพิเศษ.xlsx"

    names = collect_names(course_src, activities_src)
    print(f"พบชื่อครูจริงรวม {len(names)} ชื่อ")
    mapping = build_mapping(names)
    for real, fake in mapping.items():
        print(f"  {real} -> {fake}")

    anonymize_course(course_src, course_dst, mapping)
    anonymize_activities(activities_src, activities_dst, mapping)

    # เช็กไม่มีชื่อจริงหลงเหลือในผลลัพธ์
    remaining = set()
    for path, cols, split in ((course_dst, (0, 1), False), (activities_dst, (5,), True)):
        wb = openpyxl.load_workbook(path, data_only=True)
        for row in list(wb.active.iter_rows(values_only=True))[1:]:
            for c in cols:
                cell_vals = _split_names(row[c]) if split else ([str(row[c]).strip()] if row[c] else [])
                for v in cell_vals:
                    if v in mapping:
                        remaining.add(v)
    if remaining:
        print(f"\n[FAILED] พบชื่อจริงหลงเหลือ: {sorted(remaining)}")
        return 1

    print(f"\n[OK] สร้างไฟล์ตัวอย่างเรียบร้อย (ไม่มีชื่อจริงหลงเหลือ):")
    print(f"  {course_dst}")
    print(f"  {activities_dst}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
